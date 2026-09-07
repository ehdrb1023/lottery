from datetime import datetime
import os
import openpyxl
from rich.console import Console

console = Console()


def update_lotto_excel(
    round_num: int,
    nums: list[int],
    bonus: int,
    date_str: str = None,
    file_path: str = "data/복권_모음집.xlsx",
):
  candidates = [
      file_path,
      os.path.join(os.getcwd(), file_path),
      "data/복권_모음집.xlsx",
      "data/복권_모음집_최신판.xlsx",
      "data/복권_모음집_최신판_2.xlsx",
  ]
  target_file = None
  for cand in candidates:
    if os.path.exists(cand):
      target_file = cand
      break

  if not target_file:
    raise FileNotFoundError(f"엑셀 파일을 찾을 수 없습니다: {file_path}")

  if not date_str:
    date_str = datetime.now().strftime("%Y-%m-%d")

  sorted_nums = sorted(nums)
  wb = openpyxl.load_workbook(target_file)
  ws = wb["Sheet1"]

  # 기존 데이터의 최신 회차(Row 5) 중복 검사
  current_latest_round = ws.cell(row=5, column=6).value
  if current_latest_round and int(current_latest_round) >= round_num:
    console.print(
        f"[bold red]⚠️ 경고: 이미 {current_latest_round}회차 이상의 데이터가 존재합니다. (입력 회차: {round_num})[/bold red]"
    )
    return False

  # Row 5에 신규 행 삽입 후 데이터 기입
  ws.insert_rows(5)
  ws.cell(row=5, column=5, value=date_str)
  ws.cell(row=5, column=6, value=round_num)
  for idx, val in enumerate(sorted_nums):
    ws.cell(row=5, column=7 + idx, value=val)
  ws.cell(row=5, column=13, value=bonus)
  ws.cell(
      row=5,
      column=14,
      value="=SUMPRODUCT(COUNTIF(G5:L5,번호조회!$B$4:$B$9))",
  )

  wb.save(target_file)
  console.print(
      f"[bold green]✨ {round_num}회차 당첨 번호가 엑셀에 성공적으로 저장되었습니다![/bold green]"
  )
  console.print(
      f"   - 일자: {date_str} | 회차: {round_num}회 | 번호: {sorted_nums} + 보너스 {bonus}"
  )
  return True